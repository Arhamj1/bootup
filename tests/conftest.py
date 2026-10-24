"""A tiny synthetic corpus for the tests (not the course material)."""
import json
from pathlib import Path

import pytest

DEADLOCK_TXT = """A deadlock is a situation in which a set of processes are blocked because each process holds a resource and waits for another resource held by some other process. Four conditions must hold together for a deadlock: mutual exclusion, hold and wait, no preemption and circular wait.

The banker's algorithm avoids deadlock by granting a request only if the resulting state is safe. The safety algorithm sets Work equal to Available and Finish to false for every process. It then looks for a process whose Need is less than or equal to Work. If every process can finish, the state is safe.

A resource allocation graph draws processes and resources as nodes. A cycle in a graph with single instance resources indicates a deadlock."""

OSTEP_TXT = """Paging divides the address space into fixed-size units called pages. It avoids external fragmentation because every page frame has the same size.

A semaphore is an integer variable that is accessed only through wait and signal operations. A binary semaphore can be used as a mutex lock to protect a critical section."""

MAN_TXT = """NAME
    pthread_mutex_lock - lock a mutex

SYNOPSIS
    int pthread_mutex_lock(pthread_mutex_t *mutex);

DESCRIPTION
    The mutex object referenced by mutex is locked by calling pthread_mutex_lock. If the mutex is already locked the calling thread blocks until the mutex becomes available.
"""


@pytest.fixture
def corpus(tmp_path) -> Path:
    from pptx import Presentation
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    root = tmp_path / "raw"
    (root / "course" / "deadlocks").mkdir(parents=True)
    (root / "course" / "scheduling").mkdir(parents=True)
    (root / "open").mkdir()
    (root / "open" / "man").mkdir()
    (root / "qa_bank").mkdir()

    (root / "course" / "deadlocks" / "notes.txt").write_text(DEADLOCK_TXT)
    (root / "open" / "ostep.txt").write_text(OSTEP_TXT)
    (root / "open" / "man" / "pthread_mutex_lock.txt").write_text(MAN_TXT)

    prs = Presentation()
    layout = prs.slide_layouts[1]
    for title, body in [
        ("Round Robin", "Each process gets a time quantum. If it does not finish within the quantum it goes to the back of the ready queue. " * 3),
        ("Shortest Job First", "Picks the process with the smallest next CPU burst. It gives the optimal average waiting time but needs the burst length in advance. " * 3),
        ("Summary", "Compare the algorithms."),
        ("Note", "Practice questions."),
    ]:
        s = prs.slides.add_slide(layout)
        s.shapes.title.text = title
        s.placeholders[1].text = body
    prs.save(str(root / "course" / "scheduling" / "slides.pptx"))

    pdf = root / "course" / "paged.pdf"
    c = canvas.Canvas(str(pdf), pagesize=A4)
    bodies = [
        "Virtual memory lets a process run when only part of it is in physical memory.",
        "A page fault occurs when a process touches a page that is not in memory.",
        "The page replacement policy chooses which frame to evict on a fault.",
        "The working set model tracks the pages a process is actively using.",
        "Thrashing happens when a system spends more time paging than executing.",
    ]
    for i, body in enumerate(bodies, 1):
        c.drawString(72, 800, "Operating Systems Notes")      # running header
        c.drawString(72, 700, body)
        c.drawString(72, 40, f"Page {i}")                      # page number
        c.showPage()
    c.save()

    (root / "qa_bank" / "qa.json").write_text(json.dumps([
        {"question": "What are the four conditions for deadlock?",
         "answer": "Mutual exclusion, hold and wait, no preemption and circular wait.",
         "topic": "deadlocks", "exam": "CAT-2", "year": "2024"},
    ]))
    return root
