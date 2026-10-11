"""Turn raw course files into chunks that carry their source information.

Each source type is split by its own rule:
  slides        one chunk per slide (short consecutive slides are merged)
  prose         sentence-aligned chunks with a small overlap (PDF, DOCX, TXT, MD)
  man pages     one chunk per section (NAME, SYNOPSIS, DESCRIPTION, ...)
  qa bank       one question + answer per chunk (never split)
"""
from __future__ import annotations

import hashlib
import json
import re
from collections import Counter
from dataclasses import dataclass
from pathlib import Path

from . import config


@dataclass
class Chunk:
    id: str
    text: str
    source_file: str
    source_type: str
    topic: str
    locator: str
    exam: str = ""
    year: str = ""


@dataclass
class Unit:
    """One page / slide / paragraph as extracted, before chunking."""
    text: str
    locator: str
    title: str = ""


SUPPORTED = {".pdf", ".pptx", ".docx", ".txt", ".md", ".json"}
SENT_SPLIT = re.compile(r"(?<=[.!?])\s+")
PAGE_NUM = re.compile(r"^(page\s*)?\d{1,4}(\s*(/|of)\s*\d{1,4})?$", re.I)
MAN_HEADING = re.compile(r"^[A-Z][A-Z0-9 &/_-]{2,40}$")
MAN_FOLDERS = {"man", "man_pages", "manpages"}


# ---------------------------------------------------------------- extraction
RUN_TOGETHER = re.compile(r"\S{40,}")


def _pdf_page_text(page) -> str:
    """Default extraction, falling back to layout mode when words come out glued together
    (some slide PDFs place each word without a space character, e.g. 'readorwritedata')."""
    text = page.extract_text() or ""
    if not RUN_TOGETHER.search(text):
        return text
    try:
        laid_out = page.extract_text(extraction_mode="layout") or ""
    except Exception:
        return text
    lines = (re.sub(r"[ \t]+", " ", ln).strip() for ln in laid_out.splitlines())
    fixed = "\n".join(ln for ln in lines if ln)
    return fixed if len(RUN_TOGETHER.findall(fixed)) < len(RUN_TOGETHER.findall(text)) else text


def extract_pdf(path: Path) -> list[Unit]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    return [Unit(_pdf_page_text(page), f"p.{i}") for i, page in enumerate(reader.pages, 1)]


def extract_pptx(path: Path) -> list[Unit]:
    from pptx import Presentation

    units = []
    for i, slide in enumerate(Presentation(str(path)).slides, 1):
        title = ""
        if slide.shapes.title is not None and slide.shapes.title.has_text_frame:
            title = slide.shapes.title.text_frame.text.strip()
        parts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                t = shape.text_frame.text.strip()
                if t and t != title:
                    parts.append(t)
            elif getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    cells = [c.text.strip() for c in row.cells if c.text.strip()]
                    if cells:
                        parts.append(" | ".join(cells))
        units.append(Unit("\n".join(parts), f"slide {i}", title))
    return units


def extract_docx(path: Path) -> list[Unit]:
    from docx import Document

    return [
        Unit(p.text, f"para {i}")
        for i, p in enumerate(Document(str(path)).paragraphs, 1)
        if p.text.strip()
    ]


def extract_text(path: Path) -> list[Unit]:
    text = path.read_text(encoding="utf-8", errors="ignore")
    return [Unit(p, f"para {i}") for i, p in enumerate(re.split(r"\n\s*\n", text), 1) if p.strip()]


def split_man_sections(text: str) -> list[tuple[str, str]]:
    sections, title, buf = [], "HEADER", []
    for line in text.splitlines():
        if MAN_HEADING.match(line):
            if any(b.strip() for b in buf):
                sections.append((title, "\n".join(buf)))
            title, buf = line.strip(), []
        else:
            buf.append(line)
    if any(b.strip() for b in buf):
        sections.append((title, "\n".join(buf)))
    return sections


def load_qa(path: Path) -> list[dict]:
    data = json.loads(path.read_text(encoding="utf-8"))
    return [d for d in data if isinstance(d, dict) and d.get("question") and d.get("answer")]


# ------------------------------------------------------------------ cleaning
def clean_text(text: str) -> str:
    """Join hyphenated line breaks, drop bare page numbers, collapse whitespace."""
    text = re.sub(r"(\w)-\n([a-z])", r"\1\2", text)
    lines = [ln.strip() for ln in text.splitlines()]
    lines = [ln for ln in lines if not PAGE_NUM.match(ln)]
    text = "\n".join(lines)
    text = re.sub(r"[ \t]+", " ", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def drop_repeated_lines(units: list[Unit], min_units: int = 4, frac: float = 0.5,
                        max_len: int = 80) -> list[Unit]:
    """Remove running headers/footers: short lines present on at least `frac` of pages."""
    if len(units) < min_units:
        return units
    counts: Counter = Counter()
    for u in units:
        counts.update({ln.strip() for ln in u.text.splitlines() if ln.strip()})
    common = {ln for ln, c in counts.items() if c / len(units) >= frac and len(ln) <= max_len}
    if not common:
        return units
    return [
        Unit("\n".join(ln for ln in u.text.splitlines() if ln.strip() not in common),
             u.locator, u.title)
        for u in units
    ]


# ------------------------------------------------------------------- packing
def split_words(words: list[str], max_words: int, overlap: int) -> list[str]:
    """Hard sliding window, used only for a single over-long sentence or slide."""
    step = max(1, max_words - overlap)
    pieces = []
    for start in range(0, len(words), step):
        pieces.append(" ".join(words[start:start + max_words]))
        if start + max_words >= len(words):
            break
    return pieces


def to_atoms(units: list[Unit]) -> list[tuple[list[str], str]]:
    """Break units into sentences, remembering where each one came from."""
    atoms = []
    for u in units:
        for para in re.split(r"\n\s*\n", u.text):
            para = " ".join(para.split())
            for sent in SENT_SPLIT.split(para) if para else []:
                words = sent.split()
                if words:
                    atoms.append((words, u.locator))
    return atoms


def pack_atoms(atoms, max_words: int, overlap_words: int) -> list[tuple[str, str]]:
    """Greedily pack sentences into chunks of at most max_words.

    The last few sentences (up to overlap_words) of each chunk are repeated at
    the start of the next one so an idea split across a boundary is not lost.
    """
    out: list[tuple[str, str]] = []
    buf: list[list] = []  # [words, locator, carried_over]

    def size() -> int:
        return sum(len(a[0]) for a in buf)

    def has_fresh() -> bool:
        return any(not a[2] for a in buf)

    def emit(carry: bool) -> None:
        fresh = [a for a in buf if not a[2]]
        out.append((" ".join(w for a in buf for w in a[0]), fresh[0][1]))
        kept, total = [], 0
        if carry:
            for a in reversed(buf):
                if total + len(a[0]) > overlap_words:
                    break
                kept.insert(0, [a[0], a[1], True])
                total += len(a[0])
        buf[:] = kept

    for words, loc in atoms:
        if len(words) > max_words:
            if has_fresh():
                emit(carry=False)
            buf.clear()
            out.extend((piece, loc) for piece in split_words(words, max_words, overlap_words))
            continue
        if size() + len(words) > max_words:
            if has_fresh():
                emit(carry=True)
            if size() + len(words) > max_words:
                buf.clear()
        buf.append([words, loc, False])
    if has_fresh():
        emit(carry=False)
    return out


def pack_slides(units: list[Unit], max_words: int, overlap_words: int) -> list[tuple[str, str]]:
    out: list[tuple[str, str]] = []
    pending: list[tuple[str, str, int]] = []

    def flush() -> None:
        if not pending:
            return
        first, last = pending[0][1], pending[-1][1]
        loc = first if first == last else f"{first}-{last.split()[-1]}"
        out.append(("\n".join(p[0] for p in pending), loc))
        pending.clear()

    for u in units:
        text = f"{u.title}\n{u.text}".strip() if u.title else u.text.strip()
        n = len(text.split())
        if n == 0:
            continue
        if n > max_words:
            flush()
            out.extend((p, u.locator) for p in split_words(text.split(), max_words, overlap_words))
        elif n < config.MIN_MERGE_WORDS:
            if sum(p[2] for p in pending) + n > max_words:
                flush()
            pending.append((text, u.locator, n))
        else:
            flush()
            out.append((text, u.locator))
    flush()
    return out


def pack_sections(units: list[Unit], prefix: str, max_words: int,
                  overlap_words: int) -> list[tuple[str, str]]:
    """One chunk per man-page section; long sections are packed sentence-wise."""
    out = []
    for u in units:
        head = f"{prefix} {u.title}:".strip()
        budget = max(20, max_words - len(head.split()))
        for piece, _ in pack_atoms(to_atoms([u]), budget, overlap_words):
            out.append((f"{head} {piece}", u.title))
    return out


# --------------------------------------------------------------- per-file API
def classify(rel: Path) -> tuple[str, str]:
    """Work out (source_type, topic) from the folder layout under DATA_DIR.

    course/deadlocks/ch7.pdf          -> ("course", "deadlocks")
    open/man/processes/fork.2.txt     -> ("open", "processes")
    open/ostep.pdf                    -> ("open", "general")
    """
    folders = rel.parts[:-1]
    source_type = folders[0] if folders and folders[0] in config.SOURCE_TYPES else "course"
    named = [f for f in folders if f not in config.SOURCE_TYPES and f.lower() not in MAN_FOLDERS]
    return source_type, (named[-1] if named else "general")


def make_id(rel: Path, locator: str, text: str) -> str:
    slug = re.sub(r"[^a-z0-9]+", "-", rel.stem.lower()).strip("-")[:30] or "doc"
    h = hashlib.sha1(f"{rel.as_posix()}|{locator}|{text}".encode()).hexdigest()[:8]
    return f"{slug}-{h}"


def chunk_file(path: Path, data_dir: Path) -> list[Chunk]:
    rel = path.relative_to(data_dir)
    source_type, topic = classify(rel)
    suffix = path.suffix.lower()
    max_words = int(config.CHUNK_TOKENS * config.WORDS_PER_TOKEN)
    overlap = int(max_words * config.OVERLAP_FRACTION)
    is_man = suffix == ".man" or any(p.lower() in MAN_FOLDERS for p in rel.parts[:-1])

    def wrap(pieces, exam="", year=""):
        return [
            Chunk(make_id(rel, loc, text), text, rel.as_posix(), source_type, topic, loc, exam, year)
            for text, loc in pieces if text.strip()
        ]

    if suffix == ".json":
        if source_type != "qa_bank":
            return []
        chunks = []
        for i, qa in enumerate(load_qa(path), 1):
            text = f"Question: {qa['question'].strip()}\nAnswer: {qa['answer'].strip()}"
            chunks.append(Chunk(make_id(rel, f"qa {i}", text), text, rel.as_posix(), "qa_bank",
                                str(qa.get("topic") or topic), f"qa {i}",
                                str(qa.get("exam", "")), str(qa.get("year", ""))))
        return chunks

    if suffix in (".txt", ".md") and is_man:
        sections = split_man_sections(path.read_text(encoding="utf-8", errors="ignore"))
        units = [Unit(clean_text(body), title, title) for title, body in sections]
        return wrap(pack_sections(units, rel.stem, max_words, overlap))

    if suffix == ".pptx":
        units = drop_repeated_lines(extract_pptx(path), frac=0.6, max_len=60)
        for u in units:
            u.text = clean_text(u.text)
        return wrap(pack_slides(units, max_words, overlap))

    if suffix == ".pdf":
        units = drop_repeated_lines(extract_pdf(path))
    elif suffix == ".docx":
        units = extract_docx(path)
    elif suffix in (".txt", ".md"):
        units = extract_text(path)
    else:
        return []
    for u in units:
        u.text = clean_text(u.text)
    return wrap(pack_atoms(to_atoms(units), max_words, overlap))


def ingest_all(data_dir: Path):
    """Chunk every supported file under data_dir.

    Returns (chunks, files_used, skipped) where skipped is [(path, reason)].
    """
    chunks: list[Chunk] = []
    files: list[str] = []
    skipped: list[tuple[str, str]] = []
    for path in sorted(data_dir.rglob("*")):
        if not path.is_file() or path.name.startswith(".") or path.name.lower() == "readme.md":
            continue
        rel = path.relative_to(data_dir).as_posix()
        if path.suffix.lower() not in SUPPORTED:
            skipped.append((rel, "unsupported file type"))
            continue
        try:
            new = chunk_file(path, data_dir)
        except Exception as exc:  # keep going; report at the end
            skipped.append((rel, f"error: {exc}"))
            continue
        if not new:
            skipped.append((rel, "no extractable text (scanned PDF? it needs OCR)"))
            continue
        chunks.extend(new)
        files.append(rel)
    return chunks, files, skipped
